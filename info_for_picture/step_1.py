import re
from datetime import datetime
from openpyxl import load_workbook

# pylint: disable=R0914 too-many-locals
def forming_dict_from_collection(
        file_path: str,
        ws_title: str,
) -> dict:
    """
    В этой функции осуществляется формирование словаря из листа 'Подборки'

    :param file_path: путь до документа .xlsx
    :param ws_title: имя листа, с которого берем информацию
    :return: словарь со средствами из подборки и итоговой рекомендацией

    """
    wb = load_workbook(file_path)
    ws = wb[ws_title]

    # Определение сегодняшней даты
    today_data = datetime.today().strftime('%d.%m.%Y')

    # Поиск индексов нужных столбцов
    headers = {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}

    # Ищем строку с текущей датой и заполненным "Лучшее средство"
    target_row = None
    for row in range(2, ws.max_row + 1):
        date_cell = ws.cell(row=row, column=headers.get("Дата")).value
        if not date_cell:
            continue
        date_str = date_cell.strftime("%d.%m.%Y") \
            if isinstance(date_cell, datetime) else str(date_cell)
        best_val = ws.cell(row=row, column=headers.get("Лучшее средство")).value
        if date_str == today_data and best_val:
            target_row = row
            break

    if not target_row:
        raise ValueError("Не найдена строка с актуальной датой и заполненным 'Лучшее средство'.")

    # Находим номера средств по заголовкам вида "Средство <номер> Название"
    pattern = re.compile(r"Средство (\d+) Название")
    product_numbers = sorted(
        int(match.group(1))
        for header in headers
        if (match := pattern.fullmatch(str(header).strip()))
    )

    selection_dict = {}
    for i in product_numbers:
        name = ws.cell(row=target_row, column=headers.get(f"Средство {i} Название")).value
        if not name:
            continue
        plus = ws.cell(row=target_row, column=headers.get(f"Средство {i} ПЛЮСЫ")).value
        minus = ws.cell(row=target_row, column=headers.get(f"Средство {i} МИНУСЫ")).value
        selection_dict[name] = {
            "плюсы": plus,
            "минусы": minus,
            "лучшее средство": (name == best_val)
        }

    # Добавляем итоговую рекомендацию
    selection_dict["итоговая рекомендация"] = ws.cell(
        row=target_row, column=headers.get("Итоговая рекомендация")
    ).value

    wb.close()

    return selection_dict
