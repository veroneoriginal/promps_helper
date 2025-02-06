from openpyxl import load_workbook

def add_data_from_the_tools_page(
        ws_title: str,
        file_path: str,
        data: dict,
) -> dict:
    """Функция добавляет информацию к словарю со средствами из подборки"""

    wb = load_workbook(file_path)
    ws = wb[ws_title]

    headers = {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}

    # столбец с названием средства называется "Название"
    product_name_col_name = "Название"

    # # Собираем дополнительную информацию для каждого средства из листа "Средства"
    # # Ключ — название средства, значение — словарь с данными (цена, бренд, и т.д.)
    additional_info = {}
    for row in range(2, ws.max_row + 1):
        prod_name = ws.cell(row=row, column=headers[product_name_col_name]).value
        if not prod_name:
            continue
        # Собираем данные для всех столбцов, кроме колонки с названием продукта
        row_info = {header: ws.cell(row=row, column=col).value
                    for header, col in headers.items() if header != product_name_col_name}
        additional_info[prod_name] = row_info

    # Обновляем основной словарь, дополняя данные для каждого средства из листа "Средства"
    for prod_name in data.keys():
        # Пропускаем ключ "итоговая рекомендация"
        if prod_name == "Итоговая рекомендация":
            continue
        if prod_name in additional_info:
            data[prod_name].update(additional_info[prod_name])

    wb.close()

    return data
