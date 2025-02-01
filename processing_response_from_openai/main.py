"""В этом модуле разбираем ответ от OpanAI и записываем в таблицу"""

import json
import openpyxl


def main():
    # Пути к файлам
    excel_file = "00_base/00_Средства.xlsx"

    file_path = "prompt/history_prompt/Анализ_средств.json"

    with open(file_path, "r", encoding="utf-8") as file:
        try:
            data = json.load(file)
            print("✅ JSON корректный!")
        except json.JSONDecodeError as e:
            print(f"❌ Ошибка в JSON: {e}")

    print('Файл открыт')

    # Загружаем существующий Excel-файл
    wb = openpyxl.load_workbook(excel_file)
    # Используем лист с подборками
    ws = wb["Подборки"]

    # Ищем первую пустую строку, проверяя столбец "Лучшее средство" (H = 8)
    empty_row = None
    for row in range(2, ws.max_row + 2):
        if ws.cell(row=row, column=8).value is None:
            empty_row = row
            break

    if empty_row is None:
        empty_row = ws.max_row + 1  # Если не нашли пустую строку, добавляем в конец

    # Вставляем "Лучшее средство" (из JSON-ключа "лучшее средство")
    best_product_name = data["лучшее средство"]
    ws.cell(row=empty_row, column=8, value=best_product_name)  # Столбец "Лучшее средство"

    # Вставляем средства по порядку
    for i, item in enumerate(data["рейтинг_средств"], start=1):
        # Средство 1 начинается с колонки 9, потом +3 на каждое следующее
        col_start = 9 + (i - 1) * 3

        ws.cell(row=empty_row, column=col_start, value=item["название"])  # Название средства
        ws.cell(row=empty_row, column=col_start + 1, value="\n".join(item["плюсы"]))  # Плюсы
        ws.cell(row=empty_row, column=col_start + 2, value="\n".join(item["минусы"]))  # Минусы

    # Вставляем "Итоговую рекомендацию" (из JSON-ключа "итоговая_рекомендация")
    ws.cell(row=empty_row, column=27, value=data["итоговая_рекомендация"])

    # Сохраняем изменения
    wb.save(excel_file)

    print(f"Файл успешно обновлен: {excel_file}")


if __name__ == "__main__":
    main()
