from datetime import datetime
from pprint import pprint

from openpyxl import load_workbook

def main_step_one():
    """
    формирование словаря из листа 'Подборки'
    """

    file_path = '00_base/00_Средства.xlsx'
    wb = load_workbook(file_path)
    ws = wb['Подборки']




    today_data = datetime.today().strftime('%d.%m.%Y')
    # print(today_data)
    # print(type(today_data))

    # Поиск индексов нужных столбцов
    headers = {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}
    # print(headers)

    # Поиск строки, где значение в столбце "Дата" совпадает с сегодняшней датой
    target_row = None
    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row=row, column=headers["Дата"])
        cell_value = cell.value
        # Если в ячейке дата, преобразуем её в нужный формат
        if isinstance(cell_value, datetime):
            cell_str = cell_value.strftime("%d.%m.%Y")
        else:
            cell_str = str(cell_value)
        if cell_str == today_data:
            target_row = row
            break

    if not target_row:
        raise ValueError("Не найдена строка с сегодняшней датой.")

    # Считываем значение лучшего средства (оно должно совпадать с названием одного из средств)
    best_product_value = ws.cell(row=target_row, column=headers["Лучшее средство"]).value

    # Считываем итоговую рекомендацию
    final_rec_value = ws.cell(row=target_row, column=headers["Итоговая рекомендация"]).value


    # Формируем словарь с данными для 6 средств
    products_dict = {}
    for i in range(1, 7):
        # Считываем название, плюсы и минусы для средства i
        product_name = ws.cell(row=target_row, column=headers[f"Средство {i}"]).value
        pros = ws.cell(row=target_row, column=headers[f"Средство {i}\nПЛЮСЫ"]).value
        cons = ws.cell(row=target_row, column=headers[f"Средство {i}\nМИНУСЫ"]).value

        products_dict[product_name] = {
            "плюсы": pros,
            "минусы": cons,
            "лучшее средство": (product_name == best_product_value)
        }

    # Собираем итоговый словарь: добавляем рекомендацию

    products_dict["итоговая рекомендация"] = final_rec_value
    # Выводим результат
    # pprint(products_dict)
    return products_dict
