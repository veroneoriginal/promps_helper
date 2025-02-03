from openpyxl import load_workbook
from create_picture.step_1 import main_step_one
from pprint import pprint


def main_step_two():
    dict_from_podborki = main_step_one()

    file_path = '00_base/00_Средства.xlsx'
    wb = load_workbook(file_path)
    ws_sredstva = wb["Средства"]


    headers_sredstva = {cell.value: idx + 1 for idx, cell in enumerate(ws_sredstva[1])}

    # столбец с названием средства называется "Название"
    product_name_col_name = "Название"

    # Собираем дополнительную информацию для каждого средства из листа "Средства"
    # Ключ — название средства, значение — словарь с данными (цена, бренд, и т.д.)
    additional_info = {}
    for row in range(2, ws_sredstva.max_row + 1):
        prod_name = ws_sredstva.cell(row=row, column=headers_sredstva[product_name_col_name]).value
        if not prod_name:
            continue  # пропускаем пустые строки
        row_info = {}
        # Для каждого столбца, кроме того, где находится название средства,
        # считываем данные и записываем их в row_info
        for header, col in headers_sredstva.items():
            if header == product_name_col_name:
                continue
            row_info[header] = ws_sredstva.cell(row=row, column=col).value
        additional_info[prod_name] = row_info

    # Обновляем основной словарь, дополняя данные для каждого средства из листа "Средства"
    for prod_name in dict_from_podborki.keys():
        # Пропускаем ключ "итоговая рекомендация"
        if prod_name == "итоговая рекомендация":
            continue
        if prod_name in additional_info:
            dict_from_podborki[prod_name].update(additional_info[prod_name])
        else:
            print(f"Дополнительная информация для '{prod_name}' не найдена на листе 'Средства'.")

    print("\nИтоговый словарь после дополнения данными с листа 'Средства':")
    # pprint(dict_from_podborki)
    return dict_from_podborki

if __name__ == '__main__':
    main_step_two()